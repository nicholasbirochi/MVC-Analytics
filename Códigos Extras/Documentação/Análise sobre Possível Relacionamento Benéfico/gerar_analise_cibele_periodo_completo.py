import unicodedata
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.patches as patches
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.ticker import FuncFormatter


ANALISE_DIR = Path(__file__).resolve().parent
BASE_COMPLETA_PATH = (
    ANALISE_DIR.parent / "Analisando Irregularidades dos Registros" / "df_ronda.parquet"
)
OUTPUT_XLSX = ANALISE_DIR / "output_cibele_periodo_completo.xlsx"
OUTPUT_PDF = ANALISE_DIR / "output_cibele_periodo_completo_graficos.pdf"

CADASTRO = "4062738"
NOME = "CIBELE DE OLIVEIRA BUENO"
AREA_PREFIXO = "B-VSP"
LIMITE_SEM_PONTO_MIN = 60
HORA_CORTE_APOS_17 = 17

REGEX_PORTARIA = (
    r"\bPORTARIA\b|\bPORT\.\b|\bPORT(?=[\.\-/ ])|\bCANCELA\b|\bGUARITA\b|"
    r"\bPEDESTRES?\b|\bVISITANTE\b|\bFUNCIONARIO\b|\bVEICULOS?\b"
)
REGEX_PONTO = r"\bMARCA\b|\bPONTO\b|\bPONT\b"


def normalizar_texto(valor):
    if pd.isna(valor):
        return ""
    texto = str(valor)
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return texto.upper().strip()


def hora_to_timedelta(serie):
    hora = serie.astype("string").str.strip()
    hora = hora.replace(
        {"": "00:00:00", "<NA>": "00:00:00", "nan": "00:00:00", "NaT": "00:00:00"}
    )
    hora = hora.where(~hora.str.match(r"^\d{1,2}:\d{2}$", na=False), hora + ":00")
    return pd.to_timedelta(hora, errors="coerce")


def minutes_label(minutes):
    if pd.isna(minutes):
        return ""
    value = int(round(float(minutes)))
    return f"{value // 60:02d}:{value % 60:02d}"


def yfmt(value, pos=None):
    return minutes_label(value)


def wrap_text(ax, text, x, y, width_chars=38, line_gap=0.045, **kwargs):
    words = str(text).split()
    line = ""
    for word in words:
        test = (line + " " + word).strip()
        if len(test) > width_chars and line:
            ax.text(x, y, line, transform=ax.transAxes, **kwargs)
            y -= line_gap
            line = word
        else:
            line = test
    if line:
        ax.text(x, y, line, transform=ax.transAxes, **kwargs)
        y -= line_gap
    return y


def gaussian_smooth_line(x_values, y_values, points=260, bandwidth_days=55):
    x_values = np.asarray(x_values, dtype=float)
    y_values = np.asarray(y_values, dtype=float)
    valid = np.isfinite(x_values) & np.isfinite(y_values)
    x_values = x_values[valid]
    y_values = y_values[valid]
    if len(x_values) < 2:
        return x_values, y_values
    x_grid = np.linspace(float(np.min(x_values)), float(np.max(x_values)), points)
    distances = (x_values[:, None] - x_grid[None, :]) / float(bandwidth_days)
    weights = np.exp(-0.5 * distances**2)
    smooth = (weights * y_values[:, None]).sum(axis=0) / weights.sum(axis=0)
    return x_grid, smooth


def corrigir_data_ambigua(serie):
    data = pd.to_datetime(serie, errors="coerce").dt.normalize()
    dt = data.dt
    ambiguas = data.notna() & dt.day.le(12) & dt.month.le(12)
    invertida = pd.to_datetime(
        {"year": dt.year, "month": dt.day, "day": dt.month},
        errors="coerce",
    )
    return data.where(~ambiguas, invertida).dt.normalize()


def carregar_cibele():
    cols = [
        "cadastro",
        "nome",
        "data",
        "hora",
        "dir",
        "tipo",
        "acesso",
        "planta",
        "coletor",
        "descricao_coletor",
        "uso_acesso",
        "cargo",
        "desligado",
        "area",
        "departamento",
        "heranca",
        "cargo_especifico",
    ]
    existentes = [c for c in cols if c in pq.ParquetFile(BASE_COMPLETA_PATH).schema.names]
    df = pq.read_table(
        str(BASE_COMPLETA_PATH),
        columns=existentes,
        filters=[("cadastro", "==", CADASTRO)],
    ).to_pandas()
    if df.empty:
        raise RuntimeError(f"Nenhum registro encontrado para o cadastro {CADASTRO}.")

    df = df.rename(columns={"descricao_coletor": "descricao coletor"})
    for col in [
        "cadastro",
        "nome",
        "planta",
        "dir",
        "cargo",
        "heranca",
        "desligado",
        "area",
        "departamento",
    ]:
        if col in df.columns:
            df[col] = df[col].astype("string").str.strip()

    df = df[df["area"].astype("string").str.startswith(AREA_PREFIXO, na=False)].copy()
    if df.empty:
        raise RuntimeError(f"Nenhum registro de {NOME} encontrado em area {AREA_PREFIXO}.")

    df["data_original"] = pd.to_datetime(df["data"], errors="coerce").dt.normalize()
    df["dia_analisado"] = corrigir_data_ambigua(df["data"])
    df["_hora_td"] = hora_to_timedelta(df["hora"])
    df["_dt_ronda"] = df["dia_analisado"] + df["_hora_td"]
    df["_dir_norm"] = df["dir"].astype("string").str.upper().str.strip()
    df["planta_norm"] = df["planta"].astype("string").str.strip()
    df["home_office_999"] = df["planta_norm"].eq("999")

    desc_norm = df["descricao coletor"].map(normalizar_texto)
    df["Portaria"] = desc_norm.str.contains(REGEX_PORTARIA, regex=True, na=False)
    df["Ponto"] = desc_norm.str.contains(REGEX_PONTO, regex=True, na=False)
    df["data_registro"] = df["dia_analisado"].dt.strftime("%Y-%m-%d")
    df["hora_registro"] = df["_dt_ronda"].dt.strftime("%H:%M:%S")
    df["dia_semana_num"] = df["dia_analisado"].dt.weekday
    df["dia_semana"] = df["dia_semana_num"].map(
        {
            0: "Segunda",
            1: "Terca",
            2: "Quarta",
            3: "Quinta",
            4: "Sexta",
            5: "Sabado",
            6: "Domingo",
        }
    )
    return df.sort_values(["dia_analisado", "_dt_ronda"]).reset_index(drop=True)


def agregar_tempo(df, mascara, direcao, funcao, nome_coluna):
    filtro = mascara if direcao is None else mascara & df["_dir_norm"].eq(direcao)
    return (
        df.loc[filtro & df["_dt_ronda"].notna(), ["cadastro", "dia_analisado", "_dt_ronda"]]
        .groupby(["cadastro", "dia_analisado"], as_index=False)["_dt_ronda"]
        .agg(funcao)
        .rename(columns={"_dt_ronda": nome_coluna})
    )


def horas_positivas(fim, inicio):
    delta = fim - inicio
    return delta.dt.total_seconds().div(3600).where(delta > pd.Timedelta(0), 0).fillna(0)


def minutos_positivos(fim, inicio):
    delta = fim - inicio
    return delta.dt.total_seconds().div(60).where(delta > pd.Timedelta(0), 0).fillna(0)


def calcular_dias(df):
    meta_cols = [
        "cadastro",
        "nome",
        "area",
        "cargo",
        "heranca",
        "desligado",
        "departamento",
    ]
    meta = df[[c for c in meta_cols if c in df.columns]].drop_duplicates("cadastro")

    dias = (
        df.groupby(["cadastro", "dia_analisado"], as_index=False)
        .agg(
            nome=("nome", "first"),
            qtd_registros=("_dt_ronda", "count"),
            qtd_registros_999=("home_office_999", "sum"),
            qtd_registros_portaria=("Portaria", "sum"),
            qtd_registros_ponto=("Ponto", "sum"),
            primeira_marcacao=("_dt_ronda", "min"),
            ultima_marcacao=("_dt_ronda", "max"),
        )
        .merge(meta.drop(columns=["nome"], errors="ignore"), on="cadastro", how="left")
    )

    agregados = [
        agregar_tempo(df, df["Portaria"] & ~df["home_office_999"], "E", "min", "primeira_entrada_portaria"),
        agregar_tempo(df, df["Portaria"] & ~df["home_office_999"], "S", "max", "ultima_saida_portaria"),
        agregar_tempo(df, df["Ponto"], None, "min", "primeiro_ponto"),
        agregar_tempo(df, df["Ponto"], None, "max", "ultimo_ponto"),
    ]
    for agregado in agregados:
        dias = dias.merge(agregado, on=["cadastro", "dia_analisado"], how="left")

    dias["horas_portaria"] = horas_positivas(
        dias["ultima_saida_portaria"], dias["primeira_entrada_portaria"]
    ).round(2)
    dias["horas_ponto"] = horas_positivas(dias["ultimo_ponto"], dias["primeiro_ponto"]).round(2)
    dias["horas_servico_estimado"] = dias["horas_ponto"].where(
        dias["qtd_registros_ponto"].gt(0), dias["horas_portaria"]
    ).round(2)
    dias["ocorrencia_mais_8h_servico"] = dias["horas_servico_estimado"].gt(8)
    dias["home_office_999"] = dias["qtd_registros_999"].gt(0)

    dias["dia_semana_num"] = dias["dia_analisado"].dt.weekday
    dias["dia_semana"] = dias["dia_semana_num"].map(
        {
            0: "Segunda",
            1: "Terca",
            2: "Quarta",
            3: "Quinta",
            4: "Sexta",
            5: "Sabado",
            6: "Domingo",
        }
    )
    dias["fim_de_semana"] = dias["dia_semana_num"].ge(5)
    dias["trabalhou_fim_de_semana"] = dias["fim_de_semana"] & dias["qtd_registros"].gt(0)

    antes = minutos_positivos(dias["primeiro_ponto"], dias["primeira_entrada_portaria"])
    depois = minutos_positivos(dias["ultima_saida_portaria"], dias["ultimo_ponto"])
    dias["min_antes_do_ponto"] = antes.round(2)
    dias["min_depois_do_ponto"] = depois.round(2)
    dias["min_total_sem_ponto"] = (antes + depois).round(2)
    dias["horas_total_sem_ponto"] = (dias["min_total_sem_ponto"] / 60).round(2)
    dias["ocorrencia_sem_ponto_maior_1h"] = dias["min_total_sem_ponto"].gt(
        LIMITE_SEM_PONTO_MIN
    )

    corte_17 = dias["dia_analisado"] + pd.Timedelta(hours=HORA_CORTE_APOS_17)
    inicio_apos_17 = dias["primeira_entrada_portaria"].where(
        dias["primeira_entrada_portaria"].gt(corte_17), corte_17
    )
    apos_17 = minutos_positivos(dias["ultima_saida_portaria"], inicio_apos_17)
    dias["min_apos_17h"] = apos_17.round(2)
    dias["horas_apos_17h"] = (dias["min_apos_17h"] / 60).round(2)
    dias["ocorrencia_apos_17h"] = dias["min_apos_17h"].gt(0)

    for col in [
        "primeira_marcacao",
        "ultima_marcacao",
        "primeira_entrada_portaria",
        "ultima_saida_portaria",
        "primeiro_ponto",
        "ultimo_ponto",
    ]:
        dias[f"{col}_hora"] = dias[col].dt.strftime("%H:%M:%S")

    return dias.sort_values("dia_analisado").reset_index(drop=True)


def anexar_registros(df, dias_eventos):
    if dias_eventos.empty:
        return df.iloc[0:0].copy()
    chaves = dias_eventos[["cadastro", "dia_analisado"]].drop_duplicates()
    registros = df.merge(chaves, on=["cadastro", "dia_analisado"], how="inner")
    ordem = [
        "cadastro",
        "nome",
        "area",
        "cargo",
        "departamento",
        "heranca",
        "desligado",
        "dia_analisado",
        "data_registro",
        "hora_registro",
        "dir",
        "planta",
        "descricao coletor",
        "Portaria",
        "Ponto",
        "home_office_999",
        "tipo",
        "acesso",
    ]
    return registros[[c for c in ordem if c in registros.columns]].sort_values(
        ["dia_analisado", "hora_registro"]
    )


def dt_to_minutes(serie):
    return serie.dt.hour * 60 + serie.dt.minute + serie.dt.second / 60


def montar_estatistica(df, dias):
    dias_mais_8h = dias[dias["ocorrencia_mais_8h_servico"]]
    dias_fds = dias[dias["trabalhou_fim_de_semana"]]
    dias_apos_17 = dias[dias["ocorrencia_apos_17h"]]
    dias_sem_ponto = dias[dias["ocorrencia_sem_ponto_maior_1h"]]

    resumo = {
        "cadastro": CADASTRO,
        "nome": NOME,
        "area": ", ".join(sorted(df["area"].dropna().astype(str).unique())),
        "periodo_inicio": dias["dia_analisado"].min(),
        "periodo_fim": dias["dia_analisado"].max(),
        "linhas_total": len(df),
        "dias_avaliados": len(dias),
        "dias_com_portaria": int(dias["qtd_registros_portaria"].gt(0).sum()),
        "dias_com_ponto": int(dias["qtd_registros_ponto"].gt(0).sum()),
        "dias_mais_8h_servico": len(dias_mais_8h),
        "horas_excedentes_acima_8h": round(
            (dias_mais_8h["horas_servico_estimado"] - 8).clip(lower=0).sum(), 2
        ),
        "max_horas_servico_estimado": round(dias["horas_servico_estimado"].max(), 2),
        "dias_trabalhados_fim_de_semana": len(dias_fds),
        "dias_apos_17h": len(dias_apos_17),
        "horas_apos_17h": round(dias_apos_17["horas_apos_17h"].sum(), 2),
        "dias_sem_ponto_maior_1h": len(dias_sem_ponto),
        "horas_sem_ponto_maior_1h": round(dias_sem_ponto["horas_total_sem_ponto"].sum(), 2),
        "mediana_ultimo_ponto": minutes_label(dt_to_minutes(dias["ultimo_ponto"]).median()),
        "mediana_ultima_saida_portaria": minutes_label(
            dt_to_minutes(dias["ultima_saida_portaria"]).median()
        ),
        "base_origem": str(BASE_COMPLETA_PATH),
        "observacao_planta_999": "Planta 999 foi mantida e marcada como home_office_999.",
    }
    return pd.DataFrame([resumo])


def montar_dicionario_colunas():
    descricoes = {
        "coluna": "Nome tecnico da coluna na aba.",
        "descricao": "Explicacao curta sobre o significado da coluna.",
        "cadastro": "Matricula/cadastro da colaboradora analisada.",
        "nome": "Nome da colaboradora analisada.",
        "area": "Area registrada na base Ronda; o relatorio filtra areas iniciadas por B-VSP.",
        "cargo": "Cargo agrupado informado na base.",
        "departamento": "Departamento informado na base.",
        "heranca": "Agrupamento hierarquico herdado da base original.",
        "desligado": "Indicador da base sobre situacao de desligamento.",
        "dia_analisado": "Data consolidada do registro; datas ambiguas tiveram dia/mes corrigidos.",
        "data_original": "Data como estava armazenada no parquet antes da correcao de ambiguidades.",
        "data_registro": "Dia analisado formatado como texto no padrao AAAA-MM-DD.",
        "hora_registro": "Horario do registro depois da combinacao entre data corrigida e hora.",
        "dia_semana": "Nome do dia da semana do dia analisado.",
        "qtd_registros": "Quantidade total de registros Ronda no dia.",
        "qtd_registros_999": "Quantidade de registros com planta igual a 999 no dia.",
        "qtd_registros_portaria": "Quantidade de registros classificados como portaria no dia.",
        "qtd_registros_ponto": "Quantidade de registros classificados como ponto no dia.",
        "dir": "Direcao do registro: E para entrada, S para saida, conforme a base.",
        "planta": "Codigo da planta informado no registro.",
        "descricao coletor": "Descricao do coletor que originou o registro.",
        "Portaria": "Verdadeiro quando a descricao do coletor indica portaria, cancela, guarita ou similar.",
        "Ponto": "Verdadeiro quando a descricao do coletor indica marcacao de ponto.",
        "home_office_999": "Verdadeiro quando ha registro com planta igual a 999, tratado como home office neste relatorio.",
        "fim_de_semana": "Verdadeiro quando o dia analisado e sabado ou domingo.",
        "trabalhou_fim_de_semana": "Verdadeiro quando ha registro em sabado ou domingo.",
        "primeira_entrada_portaria_hora": "Primeiro horario de entrada em portaria no dia, desconsiderando planta 999.",
        "ultima_saida_portaria_hora": "Ultimo horario de saida em portaria no dia, desconsiderando planta 999.",
        "primeiro_ponto_hora": "Primeiro horario classificado como ponto no dia, incluindo planta 999.",
        "ultimo_ponto_hora": "Ultimo horario classificado como ponto no dia, incluindo planta 999.",
        "horas_portaria": "Intervalo, em horas, entre primeira entrada e ultima saida de portaria.",
        "horas_ponto": "Intervalo, em horas, entre primeiro e ultimo ponto do dia.",
        "horas_servico_estimado": "Estimativa de jornada: usa horas_ponto quando existe ponto; caso contrario usa horas_portaria.",
        "ocorrencia_mais_8h_servico": "Verdadeiro quando horas_servico_estimado e maior que 8 horas.",
        "min_antes_do_ponto": "Minutos entre entrada de portaria e primeiro ponto, quando o ponto ocorreu depois da entrada.",
        "min_depois_do_ponto": "Minutos entre ultimo ponto e saida de portaria, quando a saida ocorreu depois do ponto.",
        "min_total_sem_ponto": "Soma dos minutos antes e depois do ponto no dia.",
        "horas_total_sem_ponto": "min_total_sem_ponto convertido para horas.",
        "ocorrencia_sem_ponto_maior_1h": "Verdadeiro quando min_total_sem_ponto passa de 60 minutos.",
        "min_apos_17h": "Minutos entre 17h e a ultima saida de portaria, quando a saida ocorreu apos 17h.",
        "horas_apos_17h": "min_apos_17h convertido para horas.",
        "ocorrencia_apos_17h": "Verdadeiro quando houve saida de portaria apos 17h.",
        "tipo": "Tipo do registro conforme a base Ronda.",
        "acesso": "Codigo de acesso conforme a base Ronda.",
        "coletor": "Codigo do coletor conforme a base Ronda.",
        "uso_acesso": "Uso/acesso conforme a base Ronda.",
        "periodo_inicio": "Primeiro dia analisado no recorte da pessoa.",
        "periodo_fim": "Ultimo dia analisado no recorte da pessoa.",
        "linhas_total": "Total de registros brutos usados no recorte.",
        "dias_avaliados": "Total de dias com ao menos um registro no recorte.",
        "dias_com_portaria": "Quantidade de dias com ao menos um registro de portaria.",
        "dias_com_ponto": "Quantidade de dias com ao menos um registro de ponto.",
        "dias_mais_8h_servico": "Quantidade de dias com jornada estimada acima de 8 horas.",
        "horas_excedentes_acima_8h": "Soma do excedente acima de 8h nos dias com ocorrencia_mais_8h_servico.",
        "max_horas_servico_estimado": "Maior valor diario de horas_servico_estimado.",
        "dias_trabalhados_fim_de_semana": "Quantidade de dias de sabado/domingo com registros.",
        "dias_apos_17h": "Quantidade de dias com saida de portaria apos 17h.",
        "dias_sem_ponto_maior_1h": "Quantidade de dias com mais de 1h sem ponto antes/depois da portaria.",
        "horas_sem_ponto_maior_1h": "Soma de horas_total_sem_ponto nos dias acima de 1h.",
        "mediana_ultimo_ponto": "Mediana do horario do ultimo ponto diario.",
        "mediana_ultima_saida_portaria": "Mediana do horario da ultima saida diaria de portaria.",
        "base_origem": "Arquivo parquet usado como origem da analise.",
        "observacao_planta_999": "Observacao sobre o tratamento dado a planta 999.",
        "periodo": "Descricao textual do periodo considerado.",
        "filtro_area": "Filtro de area aplicado.",
        "planta_999": "Descricao textual do tratamento de planta 999.",
        "criterio_mais_8h": "Regra usada para marcar jornada acima de 8 horas.",
        "criterio_fim_de_semana": "Regra usada para marcar trabalho em fim de semana.",
        "criterio_apos_17h": "Regra usada para marcar saida apos 17h.",
        "criterio_sem_ponto": "Regra usada para marcar periodo sem ponto acima de 1h.",
        "datas_ambiguas": "Regra aplicada para datas em que dia e mes eram ambos menores ou iguais a 12.",
        "pdf_graficos": "Caminho do PDF gerado com os graficos.",
    }
    return (
        pd.DataFrame(
            [{"coluna": coluna, "descricao": descricao} for coluna, descricao in descricoes.items()]
        )
        .sort_values("coluna")
        .reset_index(drop=True)
    )


def preparar_saidas(df, dias):
    ordem_dias = [
        "cadastro",
        "nome",
        "area",
        "cargo",
        "departamento",
        "heranca",
        "desligado",
        "dia_analisado",
        "dia_semana",
        "qtd_registros",
        "qtd_registros_999",
        "qtd_registros_portaria",
        "qtd_registros_ponto",
        "home_office_999",
        "fim_de_semana",
        "trabalhou_fim_de_semana",
        "primeira_entrada_portaria_hora",
        "ultima_saida_portaria_hora",
        "primeiro_ponto_hora",
        "ultimo_ponto_hora",
        "horas_portaria",
        "horas_ponto",
        "horas_servico_estimado",
        "ocorrencia_mais_8h_servico",
        "min_antes_do_ponto",
        "min_depois_do_ponto",
        "min_total_sem_ponto",
        "horas_total_sem_ponto",
        "ocorrencia_sem_ponto_maior_1h",
        "min_apos_17h",
        "horas_apos_17h",
        "ocorrencia_apos_17h",
    ]
    registros_ordem = [
        "cadastro",
        "nome",
        "area",
        "cargo",
        "departamento",
        "heranca",
        "desligado",
        "dia_analisado",
        "data_original",
        "data_registro",
        "hora_registro",
        "dia_semana",
        "dir",
        "planta",
        "descricao coletor",
        "Portaria",
        "Ponto",
        "home_office_999",
        "tipo",
        "acesso",
        "coletor",
        "uso_acesso",
    ]
    saidas = {
        "registros_base": df[[c for c in registros_ordem if c in df.columns]].copy(),
        "dias_consolidados": dias[[c for c in ordem_dias if c in dias.columns]].copy(),
        "dias_mais_8h": dias.loc[dias["ocorrencia_mais_8h_servico"], [c for c in ordem_dias if c in dias.columns]].copy(),
        "dias_fim_semana": dias.loc[dias["trabalhou_fim_de_semana"], [c for c in ordem_dias if c in dias.columns]].copy(),
        "dias_apos_17h": dias.loc[dias["ocorrencia_apos_17h"], [c for c in ordem_dias if c in dias.columns]].copy(),
        "dias_sem_ponto": dias.loc[dias["ocorrencia_sem_ponto_maior_1h"], [c for c in ordem_dias if c in dias.columns]].copy(),
        "estatistica_geral": montar_estatistica(df, dias),
        "parametros": pd.DataFrame(
            [
                {
                    "cadastro": CADASTRO,
                    "nome": NOME,
                    "periodo": "todo o periodo disponivel para a pessoa na base",
                    "filtro_area": f"{AREA_PREFIXO}*",
                    "planta_999": "mantida como coluna/filtro home_office_999, sem analise individual",
                    "criterio_mais_8h": "horas_servico_estimado > 8; usa intervalo entre primeiro e ultimo ponto quando ha ponto, senao portaria",
                    "criterio_fim_de_semana": "dia analisado em sabado/domingo com ao menos um registro",
                    "criterio_apos_17h": "ultima saida de portaria depois de 17h",
                    "criterio_sem_ponto": "tempo antes/depois do ponto somado maior que 60 minutos",
                    "datas_ambiguas": "datas com dia e mes <= 12 foram corrigidas invertendo dia/mes, como no relatorio de saidas",
                    "pdf_graficos": str(OUTPUT_PDF),
                }
            ]
        ),
        "dicionario_colunas": montar_dicionario_colunas(),
    }
    return saidas


def escrever_excel(saidas):
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        for nome_aba, tabela in saidas.items():
            tabela.to_excel(writer, sheet_name=nome_aba, index=False)

    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(OUTPUT_XLSX)
    descricoes = dict(
        zip(
            saidas["dicionario_colunas"]["coluna"],
            saidas["dicionario_colunas"]["descricao"],
        )
    )
    for ws in wb.worksheets:
        for header_cell in ws[1]:
            descricao = descricoes.get(str(header_cell.value))
            if descricao:
                header_cell.comment = Comment(descricao, "Codex")
        ws.insert_rows(2)
        ws.row_dimensions[2].height = 42
        for desc_cell in ws[2]:
            descricao = descricoes.get(str(ws.cell(row=1, column=desc_cell.column).value), "")
            desc_cell.value = descricao
            desc_cell.font = Font(italic=True, color="5B7185", size=8)
            desc_cell.fill = PatternFill(fill_type="solid", fgColor="F8FBFD")
            desc_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = ws.dimensions
        for column_cells in ws.columns:
            col_letter = column_cells[0].column_letter
            max_len = max(
                (len(str(cell.value)) for cell in column_cells[:250] if cell.value is not None),
                default=8,
            )
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 48)
    wb.save(OUTPUT_XLSX)


def criar_pdf(dias):
    daily = dias.copy()
    daily["ultima_saida_portaria_min"] = dt_to_minutes(daily["ultima_saida_portaria"])
    daily["ultimo_ponto_min"] = dt_to_minutes(daily["ultimo_ponto"])
    daily["primeiro_ponto_min"] = dt_to_minutes(daily["primeiro_ponto"])
    daily["primeira_entrada_portaria_min"] = dt_to_minutes(daily["primeira_entrada_portaria"])

    daily_saida = daily.dropna(subset=["ultima_saida_portaria_min"]).copy()
    if daily_saida.empty:
        raise RuntimeError("Sem saidas de portaria para gerar graficos.")

    weekday_labels = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sab", "Dom"]
    daily_saida["dia_semana_curto"] = pd.Categorical(
        daily_saida["dia_semana_num"].map(dict(enumerate(weekday_labels))),
        categories=weekday_labels,
        ordered=True,
    )

    min_bin = max(
        0, int(np.floor((daily_saida["ultima_saida_portaria_min"].min() - 30) / 30) * 30)
    )
    max_bin = min(
        24 * 60,
        int(np.ceil((daily_saida["ultima_saida_portaria_min"].max() + 30) / 30) * 30),
    )
    bins = np.arange(min_bin, max_bin + 30, 30)
    labels = [f"{b // 60:02d}:{b % 60:02d}" for b in bins[:-1]]
    daily_saida["faixa_saida_portaria"] = pd.cut(
        daily_saida["ultima_saida_portaria_min"],
        bins=bins,
        labels=labels,
        right=False,
        include_lowest=True,
    )
    heat = (
        daily_saida.pivot_table(
            index="faixa_saida_portaria",
            columns="dia_semana_curto",
            values="cadastro",
            aggfunc="count",
            fill_value=0,
            observed=False,
        )
        .reindex(index=labels, columns=weekday_labels, fill_value=0)
    )
    nonzero = heat.sum(axis=1).to_numpy().nonzero()[0]
    if len(nonzero):
        heat_display = heat.iloc[max(nonzero[0] - 1, 0) : min(nonzero[-1] + 2, len(heat))]
    else:
        heat_display = heat

    x_dates = mdates.date2num(daily_saida["dia_analisado"].to_numpy(dtype="datetime64[ms]"))
    y_portaria = daily_saida["ultima_saida_portaria_min"].to_numpy(dtype=float)
    slope_day, _ = np.polyfit(x_dates, y_portaria, 1) if len(daily_saida) >= 2 else (0.0, y_portaria[0])
    slope_30d = slope_day * 30
    smooth_x, smooth_y = gaussian_smooth_line(x_dates, y_portaria)

    point_daily = daily.dropna(subset=["ultimo_ponto_min"]).copy()
    first_window = daily_saida.head(min(20, len(daily_saida)))["ultima_saida_portaria_min"].median()
    last_window = daily_saida.tail(min(20, len(daily_saida)))["ultima_saida_portaria_min"].median()
    delta_edge = last_window - first_window
    if slope_30d < -5:
        classification = "tendencia de saida mais cedo"
    elif slope_30d > 5:
        classification = "tendencia de saida mais tarde"
    else:
        classification = "sem tendencia relevante"

    period_label = f"{daily['dia_analisado'].min():%d/%m/%Y} a {daily['dia_analisado'].max():%d/%m/%Y}"
    mediana_portaria = daily_saida["ultima_saida_portaria_min"].median()
    mediana_ponto = point_daily["ultimo_ponto_min"].median()
    p10_portaria = daily_saida["ultima_saida_portaria_min"].quantile(0.10)
    p90_portaria = daily_saida["ultima_saida_portaria_min"].quantile(0.90)
    diff_mediana = mediana_portaria - mediana_ponto if not point_daily.empty else np.nan
    mode_bucket = heat.stack().idxmax() if heat.size and heat.to_numpy().max() > 0 else ("-", "-")
    mode_count = int(heat.to_numpy().max()) if heat.size else 0
    mode_txt = (
        f"{mode_bucket[1]} as {mode_bucket[0]} ({mode_count} dias)"
        if mode_bucket != ("-", "-")
        else "Sem saidas de portaria"
    )
    interp = (
        "As bolinhas representam o ultimo ponto diario; a linha mostra a tendencia suavizada "
        "da ultima saida de portaria."
    )

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "axes.titleweight": "bold",
            "axes.labelcolor": "#203348",
            "xtick.color": "#203348",
            "ytick.color": "#203348",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )
    bg = "#F8FBFD"
    ink = "#132235"
    muted = "#5B7185"
    grid = "#DCEEF7"
    blue = "#89CFF0"
    blue_dark = "#155C87"
    accent = "#F2A65A"
    cmap = LinearSegmentedColormap.from_list(
        "babyblue_clean",
        ["#FFFFFF", "#EDF9FF", "#CDEFFF", "#89CFF0", "#3E99C9"],
    )

    with PdfPages(OUTPUT_PDF) as pdf:
        fig = plt.figure(figsize=(11.69, 8.27), dpi=180, facecolor=bg)
        gs = fig.add_gridspec(
            2,
            4,
            height_ratios=[0.72, 5.2],
            width_ratios=[1, 1, 1, 1.02],
            hspace=0.16,
            wspace=0.18,
        )
        ax_title = fig.add_subplot(gs[0, :])
        ax_title.axis("off")
        ax_title.text(
            0.01,
            0.82,
            f"{NOME} - Cadastro {CADASTRO}",
            fontsize=16,
            fontweight="bold",
            color=ink,
            va="top",
        )
        ax_title.text(
            0.01,
            0.42,
            "Heatmap das ultimas saidas de portaria",
            fontsize=12.5,
            color=blue_dark,
            va="top",
        )
        ax_title.text(
            0.01,
            0.13,
            f"{period_label} | Periodo completo | Planta 999 incluida como home office",
            fontsize=8.8,
            color=muted,
            va="top",
        )

        ax = fig.add_subplot(gs[1, :3])
        arr = heat_display.to_numpy(dtype=float)
        im = ax.imshow(
            arr,
            aspect="auto",
            cmap=cmap,
            vmin=0,
            vmax=max(1, float(arr.max()) if arr.size else 1),
            origin="upper",
        )
        ax.set_xticks(np.arange(len(weekday_labels)))
        ax.set_xticklabels(weekday_labels, fontsize=9.5, fontweight="bold")
        ax.set_yticks(np.arange(len(heat_display.index)))
        ax.set_yticklabels(list(heat_display.index), fontsize=8.4)
        ax.set_xlabel("Dia da semana", fontsize=9.8, fontweight="bold", labelpad=8)
        ax.set_title("Distribuicao por faixa horaria", loc="left", fontsize=11.2, color=ink, pad=10)
        ax.set_xticks(np.arange(-0.5, len(weekday_labels), 1), minor=True)
        ax.set_yticks(np.arange(-0.5, len(heat_display.index), 1), minor=True)
        ax.grid(which="minor", color=grid, linestyle="-", linewidth=0.8)
        ax.tick_params(which="minor", bottom=False, left=False)
        for spine in ax.spines.values():
            spine.set_visible(False)
        threshold = arr.max() * 0.55 if arr.size and arr.max() else 999
        for yy in range(arr.shape[0]):
            for xx in range(arr.shape[1]):
                val = int(arr[yy, xx])
                if val:
                    ax.text(
                        xx,
                        yy,
                        str(val),
                        ha="center",
                        va="center",
                        fontsize=8.8,
                        fontweight="bold",
                        color="#FFFFFF" if val >= threshold else "#1D435A",
                    )
        cbar = fig.colorbar(im, ax=ax, fraction=0.030, pad=0.025)
        cbar.outline.set_visible(False)
        cbar.ax.tick_params(labelsize=7.8, colors=muted)
        cbar.set_label("Qtd. de dias", fontsize=8.5, color=muted)

        axp = fig.add_subplot(gs[1, 3])
        axp.axis("off")
        panel = patches.FancyBboxPatch(
            (0.02, 0.03),
            0.96,
            0.94,
            boxstyle="round,pad=0.014,rounding_size=0.018",
            transform=axp.transAxes,
            linewidth=0.9,
            edgecolor="#C8E7F7",
            facecolor="#FFFFFF",
        )
        axp.add_patch(panel)
        y0 = 0.91
        axp.text(0.08, y0, "Resumo", transform=axp.transAxes, fontsize=12, fontweight="bold", color=ink, va="top")
        y0 -= 0.085
        for label, value in [
            ("Dias avaliados", f"{len(daily)}"),
            ("Dias com portaria", f"{len(daily_saida)}"),
            ("Dias com ponto", f"{len(point_daily)}"),
            ("Mediana portaria", minutes_label(mediana_portaria)),
            ("Mediana ponto", minutes_label(mediana_ponto)),
            ("Dif. mediana", f"{diff_mediana:+.0f} min" if np.isfinite(diff_mediana) else "-"),
            ("P10-P90 portaria", f"{minutes_label(p10_portaria)}-{minutes_label(p90_portaria)}"),
        ]:
            axp.text(0.08, y0, label, transform=axp.transAxes, fontsize=8.8, color=muted, va="top")
            axp.text(
                0.92,
                y0,
                value,
                transform=axp.transAxes,
                fontsize=9.1,
                fontweight="bold",
                color=blue_dark,
                va="top",
                ha="right",
            )
            y0 -= 0.064
        y0 -= 0.032
        axp.text(0.08, y0, "Maior concentracao", transform=axp.transAxes, fontsize=9.7, fontweight="bold", color=ink, va="top")
        y0 -= 0.052
        axp.text(0.08, y0, mode_txt, transform=axp.transAxes, fontsize=9.2, color=blue_dark, va="top")
        y0 -= 0.103
        axp.text(0.08, y0, "Novos recortes", transform=axp.transAxes, fontsize=9.7, fontweight="bold", color=ink, va="top")
        y0 -= 0.052
        wrap_text(
            axp,
            f"{int(daily['ocorrencia_mais_8h_servico'].sum())} dias acima de 8h, "
            f"{int(daily['trabalhou_fim_de_semana'].sum())} dias em fim de semana. "
            "Home-office fica disponivel apenas como filtro.",
            0.08,
            y0,
            width_chars=31,
            fontsize=8.5,
            color=muted,
            va="top",
        )
        fig.subplots_adjust(left=0.055, right=0.975, top=0.94, bottom=0.08)
        pdf.savefig(fig, facecolor=bg)
        plt.close(fig)

        fig = plt.figure(figsize=(11.69, 8.27), dpi=180, facecolor=bg)
        gs = fig.add_gridspec(
            2,
            4,
            height_ratios=[0.72, 5.2],
            width_ratios=[1, 1, 1, 1.02],
            hspace=0.16,
            wspace=0.18,
        )
        ax_title = fig.add_subplot(gs[0, :])
        ax_title.axis("off")
        ax_title.text(
            0.01,
            0.82,
            f"{NOME} - Cadastro {CADASTRO}",
            fontsize=16,
            fontweight="bold",
            color=ink,
            va="top",
        )
        ax_title.text(
            0.01,
            0.42,
            "Tendencia das ultimas saidas: ponto x portaria",
            fontsize=12.5,
            color=blue_dark,
            va="top",
        )
        ax_title.text(
            0.01,
            0.13,
            "Bolinhas = ultimo ponto diario | Linha = tendencia da ultima saida de portaria",
            fontsize=8.8,
            color=muted,
            va="top",
        )

        ax = fig.add_subplot(gs[1, :3])
        ax.set_facecolor("#FFFFFF")
        ponto_handle = ax.scatter(
            point_daily["dia_analisado"],
            point_daily["ultimo_ponto_min"],
            s=24,
            color=blue,
            edgecolor=blue_dark,
            linewidth=0.35,
            alpha=0.78,
            label="Ultimo ponto diario",
        )
        tendencia_handle = ax.plot(
            smooth_x,
            smooth_y,
            color=blue_dark,
            linewidth=3.0,
            solid_capstyle="round",
            solid_joinstyle="round",
            label="Tendencia da saida de portaria",
        )[0]
        home_days = point_daily[point_daily["home_office_999"]]
        home_handle = None
        if not home_days.empty:
            home_handle = ax.scatter(
                home_days["dia_analisado"],
                home_days["ultimo_ponto_min"],
                s=42,
                facecolors="none",
                edgecolors=accent,
                linewidth=1.1,
                alpha=0.9,
                label="Dia com Home-office",
            )
        ax.set_title("Evolucao temporal", loc="left", fontsize=11.2, color=ink, pad=10)
        ax.set_ylabel("Horario", fontsize=9.8, fontweight="bold", labelpad=8)
        ax.set_xlabel("Data", fontsize=9.8, fontweight="bold", labelpad=8)
        ax.yaxis.set_major_formatter(FuncFormatter(yfmt))
        ax.xaxis.set_major_locator(mdates.MonthLocator(interval=2))
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%b/%y"))
        plt.setp(ax.get_xticklabels(), rotation=0, ha="center", fontsize=8.5)
        all_y = np.concatenate(
            [
                point_daily["ultimo_ponto_min"].dropna().to_numpy(dtype=float),
                daily_saida["ultima_saida_portaria_min"].dropna().to_numpy(dtype=float),
            ]
        )
        min_y = max(6 * 60, np.floor((np.nanmin(all_y) - 45) / 30) * 30)
        max_y = min(23 * 60, np.ceil((np.nanmax(all_y) + 45) / 30) * 30)
        ax.set_ylim(max_y, min_y)
        ax.set_yticks(np.arange(min_y, max_y + 1, 30))
        ax.grid(True, axis="both", color=grid, linewidth=0.8)
        for spine in ax.spines.values():
            spine.set_visible(False)
        legend_handles = [ponto_handle]
        if home_handle is not None:
            legend_handles.append(home_handle)
        legend_handles.append(tendencia_handle)
        legend = ax.legend(
            handles=legend_handles,
            loc="upper left",
            frameon=True,
            facecolor="#FFFFFF",
            edgecolor=grid,
            fontsize=8.4,
        )
        legend.get_frame().set_linewidth(0.7)

        axp = fig.add_subplot(gs[1, 3])
        axp.axis("off")
        panel = patches.FancyBboxPatch(
            (0.02, 0.03),
            0.96,
            0.94,
            boxstyle="round,pad=0.014,rounding_size=0.018",
            transform=axp.transAxes,
            linewidth=0.9,
            edgecolor="#C8E7F7",
            facecolor="#FFFFFF",
        )
        axp.add_patch(panel)
        y0 = 0.91
        axp.text(0.08, y0, "Tendencia", transform=axp.transAxes, fontsize=12, fontweight="bold", color=ink, va="top")
        y0 -= 0.088
        axp.text(0.08, y0, "Conclusao", transform=axp.transAxes, fontsize=8.8, color=muted, va="top")
        y0 -= 0.043
        axp.text(
            0.08,
            y0,
            classification,
            transform=axp.transAxes,
            fontsize=9.2,
            fontweight="bold",
            color=blue_dark,
            va="top",
        )
        y0 -= 0.078
        for label, value in [
            ("Inclinacao", f"{slope_30d:+.1f} min / 30 dias"),
            ("Mediana inicial", minutes_label(first_window)),
            ("Mediana final", minutes_label(last_window)),
            ("Diferenca", f"{delta_edge:+.0f} min"),
            ("Dias ponto", f"{len(point_daily)}"),
            ("Dias portaria", f"{len(daily_saida)}"),
        ]:
            axp.text(0.08, y0, label, transform=axp.transAxes, fontsize=8.8, color=muted, va="top")
            axp.text(
                0.92,
                y0,
                value,
                transform=axp.transAxes,
                fontsize=8.9,
                fontweight="bold",
                color=blue_dark,
                va="top",
                ha="right",
            )
            y0 -= 0.066
        y0 -= 0.045
        axp.text(0.08, y0, "Interpretacao", transform=axp.transAxes, fontsize=9.7, fontweight="bold", color=ink, va="top")
        y0 -= 0.055
        wrap_text(axp, interp, 0.08, y0, width_chars=32, fontsize=8.6, color=muted, va="top")
        fig.subplots_adjust(left=0.055, right=0.975, top=0.94, bottom=0.08)
        pdf.savefig(fig, facecolor=bg)
        plt.close(fig)


def main():
    df = carregar_cibele()
    dias = calcular_dias(df)
    saidas = preparar_saidas(df, dias)
    escrever_excel(saidas)
    criar_pdf(dias)

    resumo = saidas["estatistica_geral"].iloc[0].to_dict()
    print(f"Excel: {OUTPUT_XLSX}")
    print(f"PDF: {OUTPUT_PDF}")
    print("Resumo:")
    for chave in [
        "linhas_total",
        "dias_avaliados",
        "dias_mais_8h_servico",
        "dias_trabalhados_fim_de_semana",
        "dias_apos_17h",
        "horas_apos_17h",
    ]:
        print(f"- {chave}: {resumo[chave]}")


if __name__ == "__main__":
    main()
